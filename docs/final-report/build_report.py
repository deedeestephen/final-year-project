#!/usr/bin/env python3
"""Build the final-year report as a Word document from the Markdown parts in this folder.

    python build_report.py            # writes PCa-mHealth-final-report.docx
    python build_report.py --check    # also reads the file back and prints a summary

The Markdown files are the source of truth; edit them, then run this script again.
It understands the small subset of Markdown the report uses:

- ``#`` to ``####`` headings (real Word heading styles, so Word can build the contents);
- paragraphs with **bold**, *italic*, `code` and [links](url) (the link text is kept);
- ``- `` bullet lists and ``1. `` numbered lists;
- pipe tables, with an optional caption paragraph "Table N.N: ..." just above them;
- an image on its own line, ``![Figure N.N: caption](path.png)``: the picture is inserted
  with the caption below it; a line holding several images becomes a picture grid, and the
  paragraph after it, "Figure N.N: ...", is its caption;
- fenced code blocks;
- markers in HTML comments: ``<!-- titlepage -->`` ... ``<!-- /titlepage -->`` and
  ``<!-- toc -->`` (contents, list of figures and list of tables). Other comments are dropped.

Placeholders in square brackets that start with capital letters, such as [STUDENT NAME],
are highlighted in yellow so that none is missed.
"""

from __future__ import annotations

import argparse
import math
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
PARTS = [
    "00-front-matter.md",
    "01-introduction.md",
    "02-literature-review.md",
    "03-system-analysis-and-design.md",
    "04-implementation-and-testing.md",
    "05-conclusion-and-recommendations.md",
    "06-references.md",
    "07-appendices.md",
]
OUTPUT = HERE / "PCa-mHealth-final-report.docx"
TITLE = (
    "An Artificial Intelligence-Driven Mobile Health Platform for Enhanced Prostate Cancer "
    "Diagnosis and Detection within the Zambian Healthcare System"
)

# Page set-up: A4, 1-inch margins. Text area in inches for each orientation.
PAGE_W, PAGE_H, MARGIN = 8.27, 11.69, 1.0
TEXT_W_PORTRAIT = PAGE_W - 2 * MARGIN  # 6.27
TEXT_H_PORTRAIT = PAGE_H - 2 * MARGIN  # 9.69
TEXT_W_LANDSCAPE = PAGE_H - 2 * MARGIN
TEXT_H_LANDSCAPE = PAGE_W - 2 * MARGIN
CAPTION_ROOM = 1.2  # inches kept free for the caption and spacing under a figure
WIDE_DIAGRAM_RATIO = 2.3  # wider diagrams go on a landscape page

BODY_FONT = "Times New Roman"
CODE_FONT = "Courier New"
BODY_PT = 12
TABLE_PT = 10

FIGURE_CAPTION = "Figure Caption"
TABLE_CAPTION = "Table Caption"
CONTENTS_TITLE = "Contents Title"
REFERENCE = "Reference"
CODE_BLOCK = "Code Block"

PLACEHOLDER = re.compile(r"\[[A-Z][A-Z0-9 ,'&/\-]*[A-Z](?::[^\]]*)?\]")
NOT_PLACEHOLDERS = {"[WHO]"}
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\s](?:[^*]*?[^*\s])?\*|\[[^\]]+\]\([^)]+\))")
IMAGE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
IMAGE_LINE = re.compile(r"^(?:!\[[^\]]*\]\([^)]+\)\s*)+$")
CAPTION_LINE = re.compile(r"^(Figure|Table) [0-9A-Z]+\.\d+: ")
NUMBERED = re.compile(r"^(\d+)\.\s+(.*)$")


# Elements that must come after a given element inside its parent (ECMA-376 order). Word refuses
# some documents whose elements are out of order, so new elements are inserted before these.
AFTER_PGNUMTYPE = ("cols", "formProt", "vAlign", "noEndnote", "titlePg", "textDirection", "bidi",
                   "rtlGutter", "docGrid", "printerSettings", "sectPrChange")
AFTER_PPR_SHD = ("tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct",
                 "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
                 "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
                 "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
                 "pPrChange")
AFTER_TCPR_SHD = ("noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark", "headers",
                  "cellIns", "cellDel", "cellMerge", "tcPrChange")
AFTER_CANTSPLIT = ("trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden", "ins", "del", "trPrChange")
AFTER_TBLHEADER = ("tblCellSpacing", "jc", "hidden", "ins", "del", "trPrChange")
AFTER_UPDATEFIELDS = ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
                      "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                      "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown",
                      "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags",
                      "decimalSymbol", "listSeparator")


def insert_in_order(parent, child, successors) -> None:
    """Insert child before the first existing element that the schema puts after it."""
    names = {qn(f"w:{name}") for name in successors}
    for index, existing in enumerate(parent):
        if existing.tag in names:
            parent.insert(index, child)
            return
    parent.append(child)


# --------------------------------------------------------------------------- styles


def set_run_font(rpr_owner, name: str) -> None:
    """Set a font on a style or run for every script, and drop theme fonts that would override it."""
    rpr = rpr_owner.element.get_or_add_rPr() if hasattr(rpr_owner, "element") else rpr_owner
    fonts = rpr.get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)


def paragraph_style(doc, name: str, base: str = "Normal"):
    styles = doc.styles
    try:
        return styles[name]
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles[base]
        style.quick_style = True
        return style


def configure_styles(doc) -> None:
    normal = doc.styles["Normal"]
    normal.font.size = Pt(BODY_PT)
    set_run_font(normal, BODY_FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)

    sizes = {1: 16, 2: 14, 3: 12, 4: 12}
    for level, size in sizes.items():
        style = doc.styles[f"Heading {level}"]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = level == 4
        style.font.color.rgb = RGBColor(0, 0, 0)
        set_run_font(style, BODY_FONT)
        hpf = style.paragraph_format
        hpf.keep_with_next = True
        hpf.line_spacing = 1.15
        hpf.space_before = Pt({1: 0, 2: 18, 3: 12, 4: 10}[level])
        hpf.space_after = Pt({1: 18, 2: 8, 3: 6, 4: 6}[level])
        hpf.page_break_before = level == 1

    for name in ("List Bullet", "List Bullet 2"):
        style = doc.styles[name]
        style.font.size = Pt(BODY_PT)
        set_run_font(style, BODY_FONT)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(3)

    numbered = paragraph_style(doc, "List Numbered Manual")
    numbered.paragraph_format.left_indent = Inches(0.35)
    numbered.paragraph_format.first_line_indent = Inches(-0.35)
    numbered.paragraph_format.tab_stops.add_tab_stop(Inches(0.35))
    numbered.paragraph_format.space_after = Pt(3)

    for name, align in ((FIGURE_CAPTION, WD_ALIGN_PARAGRAPH.CENTER), (TABLE_CAPTION, WD_ALIGN_PARAGRAPH.LEFT)):
        style = paragraph_style(doc, name)
        style.font.size = Pt(11)
        style.font.italic = True
        style.paragraph_format.alignment = align
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_before = Pt(4 if name == FIGURE_CAPTION else 10)
        style.paragraph_format.space_after = Pt(12 if name == FIGURE_CAPTION else 4)
        style.paragraph_format.keep_with_next = name == TABLE_CAPTION

    contents = paragraph_style(doc, CONTENTS_TITLE)
    contents.font.size = Pt(16)
    contents.font.bold = True
    contents.paragraph_format.space_after = Pt(18)
    contents.paragraph_format.page_break_before = True

    reference = paragraph_style(doc, REFERENCE)
    reference.paragraph_format.left_indent = Inches(0.5)
    reference.paragraph_format.first_line_indent = Inches(-0.5)
    reference.paragraph_format.space_after = Pt(8)
    reference.paragraph_format.line_spacing = 1.5

    code = paragraph_style(doc, CODE_BLOCK)
    code.font.size = Pt(9.5)
    set_run_font(code, CODE_FONT)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.left_indent = Inches(0.3)

    cell = paragraph_style(doc, "Table Text")
    cell.font.size = Pt(TABLE_PT)
    cell.paragraph_format.line_spacing = 1.0
    cell.paragraph_format.space_after = Pt(0)

    grid_label = paragraph_style(doc, "Grid Label", base="Table Text")
    grid_label.font.size = Pt(9)
    grid_label.font.italic = True
    grid_label.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


# --------------------------------------------------------------------------- page set-up and fields


def set_page(section, landscape: bool) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    w, h = (PAGE_H, PAGE_W) if landscape else (PAGE_W, PAGE_H)
    section.page_width, section.page_height = Inches(w), Inches(h)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(MARGIN))


def page_numbering(section, fmt: str, start: int | None) -> None:
    sect_pr = section._sectPr
    pg = sect_pr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        insert_in_order(sect_pr, pg, AFTER_PGNUMTYPE)
    pg.set(qn("w:fmt"), fmt)
    if start is None:
        if pg.get(qn("w:start")) is not None:
            del pg.attrib[qn("w:start")]
    else:
        pg.set(qn("w:start"), str(start))


def add_field(paragraph, instruction: str, placeholder: str) -> None:
    def fld(kind: str):
        run = paragraph.add_run()
        char = OxmlElement("w:fldChar")
        char.set(qn("w:fldCharType"), kind)
        if kind == "begin":
            char.set(qn("w:dirty"), "true")
        run._r.append(char)

    fld("begin")
    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    instr_run._r.append(instr)
    fld("separate")
    paragraph.add_run(placeholder)
    fld("end")


def add_page_number_footer(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(paragraph, "PAGE", "1")
    for run in paragraph.runs:
        run.font.size = Pt(11)
        set_run_font(run._r.get_or_add_rPr(), BODY_FONT)


def ask_word_to_update_fields(doc) -> None:
    """Word then offers to update the contents, the lists and the page numbers when the file opens."""
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        insert_in_order(settings, update, AFTER_UPDATEFIELDS)
    update.set(qn("w:val"), "true")


# --------------------------------------------------------------------------- inline text


def add_text(paragraph, text: str, bold=False, italic=False, size: float | None = None) -> None:
    """Add text, highlighting placeholders such as [STUDENT NAME]."""
    pos = 0
    for match in PLACEHOLDER.finditer(text):
        if match.group(0) in NOT_PLACEHOLDERS:
            continue
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            style_run(run, bold, italic, size)
        run = paragraph.add_run(match.group(0))
        style_run(run, bold, italic, size)
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        pos = match.end()
    if pos < len(text):
        style_run(paragraph.add_run(text[pos:]), bold, italic, size)


def style_run(run, bold, italic, size) -> None:
    run.bold = bold or None
    run.italic = italic or None
    if size:
        run.font.size = Pt(size)


def add_inline(paragraph, text: str, bold=False, italic=False, size: float | None = None) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            add_inline(paragraph, part[2:-2], True, italic, size)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run._r.get_or_add_rPr(), CODE_FONT)
            run.font.size = Pt((size or BODY_PT) - 1.5)
            style_run(run, bold, italic, None)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            add_inline(paragraph, part[1:-1], bold, True, size)
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            add_inline(paragraph, part[1 : part.index("](")], bold, italic, size)
        else:
            add_text(paragraph, part, bold, italic, size)


# --------------------------------------------------------------------------- blocks


def parse_blocks(text: str):
    """Yield (kind, payload) blocks from the Markdown text."""
    text = text.replace("\r\n", "\n")
    text = re.sub(r"<!--\s*(/?titlepage|toc)\s*-->", r"@@\1@@", text)
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    lines = text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("@@") and stripped.endswith("@@"):
            yield "marker", stripped.strip("@")
            i += 1
            continue
        if stripped.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            i += 1
            yield "code", code
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            yield "heading", (len(heading.group(1)), heading.group(2).strip())
            i += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            yield "table", rows
            continue
        if IMAGE_LINE.match(stripped):
            yield "images", IMAGE.findall(stripped)
            i += 1
            continue
        if re.match(r"^\s*[-*]\s+", line):
            indent = len(line) - len(line.lstrip())
            yield "bullet", (indent, re.sub(r"^\s*[-*]\s+", "", line))
            i += 1
            continue
        number = NUMBERED.match(stripped)
        if number:
            yield "number", (number.group(1), number.group(2))
            i += 1
            continue
        paragraph = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||!\[|[-*]\s|\d+\.\s|```|@@)", lines[i].strip()):
            paragraph.append(lines[i].strip())
            i += 1
        yield "paragraph", " ".join(paragraph)


def split_row(row: str) -> list[str]:
    row = row.strip()
    if row.startswith("|"):
        row = row[1:]
    if row.endswith("|"):
        row = row[:-1]
    return [cell.strip() for cell in row.split("|")]


class ReportBuilder:
    def __init__(self) -> None:
        self.doc = Document()
        configure_styles(self.doc)
        section = self.doc.sections[0]
        set_page(section, landscape=False)
        page_numbering(section, "lowerRoman", 1)
        section.different_first_page_header_footer = True  # no number on the title page
        add_page_number_footer(section)
        self.landscape = False
        self.part = ""
        self.in_titlepage = False
        self.pending_caption_for_grid = False
        self.counts = {"figures": 0, "tables": 0, "front_tables": 0, "grids": 0, "landscape_pages": 0}
        self.figure_numbers: list[str] = []
        self.table_numbers: list[str] = []
        core = self.doc.core_properties
        core.title = TITLE
        core.subject = "Final-year project report, ZCAS University (draft)"
        core.author = "[STUDENT NAME]"
        core.comments = "Draft prepared with AI assistance from the project documentation; see README.md."

    # ----- sections

    def new_section(self, landscape: bool, restart_numbers: bool = False, fmt: str = "decimal") -> None:
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        set_page(section, landscape)
        page_numbering(section, fmt, 1 if restart_numbers else None)
        section.different_first_page_header_footer = False
        self.landscape = landscape

    # ----- text blocks

    def heading(self, level: int, text: str) -> None:
        if self.in_titlepage:
            self.title_line(text)
            return
        paragraph = self.doc.add_paragraph(style=f"Heading {level}")
        add_inline(paragraph, text)

    def title_line(self, text: str) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        is_title = text.startswith("**") and text.endswith("**")
        paragraph.paragraph_format.space_before = Pt(48 if is_title else 6)
        paragraph.paragraph_format.space_after = Pt(36 if is_title else 6)
        if is_title:
            add_inline(paragraph, text.strip("*"), bold=True, size=18)
        elif text.isupper():
            add_inline(paragraph, text, bold=True, size=14)
        else:
            add_inline(paragraph, text)

    def paragraph(self, text: str) -> None:
        if self.in_titlepage:
            self.title_line(text)
            return
        caption = CAPTION_LINE.match(text)
        if caption and caption.group(1) == "Table":
            paragraph = self.doc.add_paragraph(style=TABLE_CAPTION)
            add_inline(paragraph, text)
            self.table_numbers.append(text.split(":")[0])
            return
        if caption and caption.group(1) == "Figure":
            paragraph = self.doc.add_paragraph(style=FIGURE_CAPTION)
            add_inline(paragraph, text)
            if self.pending_caption_for_grid:
                self.figure_numbers.append(text.split(":")[0])
                self.pending_caption_for_grid = False
            return
        style = REFERENCE if self.part.startswith("06") else None
        paragraph = self.doc.add_paragraph(style=style)
        add_inline(paragraph, text)

    def bullet(self, indent: int, text: str) -> None:
        paragraph = self.doc.add_paragraph(style="List Bullet 2" if indent >= 2 else "List Bullet")
        add_inline(paragraph, text)

    def number(self, label: str, text: str) -> None:
        paragraph = self.doc.add_paragraph(style="List Numbered Manual")
        paragraph.add_run(f"{label}.\t")
        add_inline(paragraph, text)

    def code(self, lines: list[str]) -> None:
        for n, line in enumerate(lines):
            paragraph = self.doc.add_paragraph(style=CODE_BLOCK)
            paragraph.add_run(line or " ")
            shade(paragraph._p.get_or_add_pPr(), "F2F2F2", AFTER_PPR_SHD)
            if n == len(lines) - 1:
                paragraph.paragraph_format.space_after = Pt(8)

    def marker(self, name: str) -> None:
        if name == "titlepage":
            self.in_titlepage = True
        elif name == "/titlepage":
            self.in_titlepage = False
        elif name == "toc":
            self.contents()

    def contents(self) -> None:
        for title, instruction in (
            ("Table of Contents", 'TOC \\o "1-3" \\h \\z \\u'),
            ("List of Figures", f'TOC \\h \\z \\t "{FIGURE_CAPTION},1"'),
            ("List of Tables", f'TOC \\h \\z \\t "{TABLE_CAPTION},1"'),
        ):
            self.doc.add_paragraph(title, style=CONTENTS_TITLE)
            paragraph = self.doc.add_paragraph()
            add_field(paragraph, instruction, "Right-click here and choose Update Field (or press F9) to build this list.")

    # ----- tables

    def table(self, rows: list[str]) -> None:
        cells = [split_row(row) for row in rows if not re.match(r"^\|?\s*:?-{3,}", row.strip())]
        if not cells:
            return
        columns = max(len(row) for row in cells)
        cells = [row + [""] * (columns - len(row)) for row in cells]
        table = self.doc.add_table(rows=len(cells), cols=columns)
        table.style = self.doc.styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = column_widths(cells, TEXT_W_LANDSCAPE if self.landscape else TEXT_W_PORTRAIT)
        for r, row in enumerate(cells):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.width = Inches(widths[c])
                paragraph = cell.paragraphs[0]
                paragraph.style = self.doc.styles["Table Text"]
                add_inline(paragraph, value, bold=(r == 0), size=TABLE_PT)
                if r == 0:
                    shade(cell._tc.get_or_add_tcPr(), "D9E2F3", AFTER_TCPR_SHD)
            no_split(table.rows[r])
        repeat_header(table.rows[0])
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)
        spacer.paragraph_format.line_spacing = 1.0
        self.counts["front_tables" if self.part.startswith("00") else "tables"] += 1

    # ----- figures

    def images(self, found: list[tuple[str, str]]) -> None:
        if len(found) == 1:
            self.figure(found[0][0], found[0][1])
        else:
            self.grid(found)

    def resolve(self, path: str) -> Path:
        image = (HERE / path).resolve()
        if not image.is_file():
            raise FileNotFoundError(f"{self.part}: picture not found: {path}")
        return image

    def figure(self, caption: str, path: str) -> None:
        image = self.resolve(path)
        info = DocxImage.from_file(str(image))
        ratio = info.px_width / info.px_height
        wide = "report/figures" in path.replace("\\", "/") and ratio >= WIDE_DIAGRAM_RATIO
        was_landscape = self.landscape
        if wide and not was_landscape:
            self.new_section(landscape=True)
            self.counts["landscape_pages"] += 1
        text_w = TEXT_W_LANDSCAPE if self.landscape else TEXT_W_PORTRAIT
        text_h = (TEXT_H_LANDSCAPE if self.landscape else TEXT_H_PORTRAIT) - CAPTION_ROOM
        if ratio < 0.6 and "report/img/app-" in path.replace("\\", "/"):
            text_h = min(text_h, 5.5)  # a single phone screen: keep it at a readable, not page-filling, size
        min_dpi = 200 if "report/figures" in path.replace("\\", "/") else 150
        width = min(text_w, text_h * ratio, info.px_width / min_dpi)
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.add_run().add_picture(str(image), width=Inches(width))
        self.counts["figures"] += 1
        if CAPTION_LINE.match(caption):
            captioned = self.doc.add_paragraph(style=FIGURE_CAPTION)
            add_inline(captioned, caption)
            self.figure_numbers.append(caption.split(":")[0])
        if wide and not was_landscape:
            self.new_section(landscape=False)

    def grid(self, found: list[tuple[str, str]]) -> None:
        paths = [self.resolve(path) for _, path in found]
        infos = [DocxImage.from_file(str(p)) for p in paths]
        ratios = [i.px_width / i.px_height for i in infos]
        text_w = TEXT_W_LANDSCAPE if self.landscape else TEXT_W_PORTRAIT
        phones = all(r < 0.6 for r in ratios)
        if phones:
            columns = len(found) if len(found) <= 4 else math.ceil(len(found) / 2)
        else:
            columns = 1
        rows = math.ceil(len(found) / columns)
        table = self.doc.add_table(rows=rows, cols=columns)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell_w = text_w / columns
        max_h = (TEXT_H_PORTRAIT - CAPTION_ROOM - 0.4 * rows) / rows
        for n, ((label, _), image, ratio) in enumerate(zip(found, paths, ratios)):
            cell = table.cell(n // columns, n % columns)
            cell.width = Inches(cell_w)
            width = min(cell_w - 0.15, max_h * ratio)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(image), width=Inches(width))
            if label:
                label_paragraph = cell.add_paragraph(style="Grid Label")
                add_inline(label_paragraph, label, italic=True, size=9)
        for row in table.rows:
            no_split(row)
        self.counts["figures"] += 1
        self.counts["grids"] += 1
        self.pending_caption_for_grid = True

    # ----- whole document

    def build(self) -> Path:
        for name in PARTS:
            self.part = name
            if name == "01-introduction.md":
                # the main text: Arabic page numbers from 1
                self.new_section(landscape=False, restart_numbers=True)
            text = (HERE / name).read_text(encoding="utf-8")
            for kind, payload in parse_blocks(text):
                if kind == "heading":
                    self.heading(*payload)
                elif kind == "paragraph":
                    self.paragraph(payload)
                elif kind == "bullet":
                    self.bullet(*payload)
                elif kind == "number":
                    self.number(*payload)
                elif kind == "table":
                    self.table(payload)
                elif kind == "images":
                    self.images(payload)
                elif kind == "code":
                    self.code(payload)
                elif kind == "marker":
                    self.marker(payload)
            if self.pending_caption_for_grid:
                raise ValueError(f"{name}: a picture grid has no 'Figure N.N:' caption after it")
        ask_word_to_update_fields(self.doc)
        self.doc.save(OUTPUT)
        return OUTPUT


# --------------------------------------------------------------------------- table helpers


def column_widths(cells: list[list[str]], total: float) -> list[float]:
    columns = len(cells[0])
    weights = []
    for c in range(columns):
        lengths = [len(re.sub(r"[*`]", "", row[c])) for row in cells]
        longest = max(lengths)
        average = sum(lengths) / len(lengths)
        weights.append(max(6.0, min(60.0, 0.5 * longest + 0.5 * average)))
    raw = [total * w / sum(weights) for w in weights]
    minimum = min(0.75, total / columns)
    widths = [max(minimum, w) for w in raw]
    scale = total / sum(widths)
    return [w * scale for w in widths]


def shade(properties, fill: str, successors) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    insert_in_order(properties, shading, successors)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    insert_in_order(tr_pr, header, AFTER_TBLHEADER)


def no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    insert_in_order(tr_pr, cant_split, AFTER_CANTSPLIT)


# --------------------------------------------------------------------------- read-back check


def check(path: Path, builder: ReportBuilder) -> int:
    """Open the written file again and report what is in it."""
    doc = Document(str(path))
    headings = {}
    for paragraph in doc.paragraphs:
        name = paragraph.style.name if paragraph.style is not None else ""
        if name.startswith("Heading "):
            headings[name] = headings.get(name, 0) + 1
    figure_captions = sum(1 for p in doc.paragraphs if p.style is not None and p.style.name == FIGURE_CAPTION)
    table_captions = sum(1 for p in doc.paragraphs if p.style is not None and p.style.name == TABLE_CAPTION)
    pictures = len(doc.inline_shapes)
    texts = [p.text for p in doc.paragraphs]
    texts += [p.text for t in doc.tables for row in t.rows for cell in row.cells for p in cell.paragraphs]
    placeholders = sorted(
        {m.group(0).split(":")[0] + (": ...]" if ":" in m.group(0) else "") for t in texts for m in PLACEHOLDER.finditer(t)}
        - NOT_PLACEHOLDERS
    )
    sections = doc.sections
    print(f"Opened {path.name}: {path.stat().st_size / 1e6:.1f} MB")
    print(f"  sections: {len(sections)} ({sum(1 for s in sections if s.orientation == WD_ORIENT.LANDSCAPE)} landscape)")
    print(f"  headings: " + ", ".join(f"{k}: {v}" for k, v in sorted(headings.items())))
    print(f"  pictures embedded: {pictures}; figures placed: {builder.counts['figures']} (of which grids: {builder.counts['grids']})")
    print(f"  figure captions: {figure_captions}; table captions: {table_captions}")
    print(f"  data tables: {builder.counts['tables']} in the chapters and appendices, {builder.counts['front_tables']} in the front matter"
          f" (Word tables in the file, including picture grids: {len(doc.tables)})")
    numbers = builder.figure_numbers
    duplicates = sorted({n for n in numbers if numbers.count(n) > 1})
    print(f"  figure numbers: {numbers[0]} to {numbers[-1]} ({len(numbers)}; duplicates: {duplicates or 'none'})")
    print(f"  table numbers: {len(builder.table_numbers)} captions")
    print(f"  placeholders still to fill: {', '.join(placeholders)}")
    problems = 0
    if figure_captions != builder.counts["figures"]:
        print("  PROBLEM: not every figure has a caption")
        problems += 1
    if table_captions != builder.counts["tables"]:
        print("  PROBLEM: not every table has a caption")
        problems += 1
    if duplicates:
        print("  PROBLEM: duplicate figure numbers")
        problems += 1
    return problems


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--check", action="store_true", help="read the file back and print a summary")
    args = parser.parse_args()
    builder = ReportBuilder()
    path = builder.build()
    print(f"Wrote {path}")
    if args.check:
        return 1 if check(path, builder) else 0
    return 0


if __name__ == "__main__":
    sys.exit(main())
